from __future__ import annotations

import os
from datetime import datetime
from io import BytesIO
from pathlib import Path
from threading import Lock

from openpyxl import Workbook, load_workbook

from app.infrastructure.supabase_client import get_supabase_client
from app.repositories.convocatoria_repository import ConvocatoriaRepository
from app.schemas.student import RegistroEstudianteRequest


class StudentRepository:
    """Repositorio híbrido: Supabase en producción y Excel como fallback local."""

    _lock = Lock()
    _headers = ["Fecha", "Nombre", "CI", "RU", "Correo", "Celular"]
    _incident_headers = ["Fecha", "Nombre", "CI", "RU", "Archivo", "Motivo"]

    def __init__(self) -> None:
        self.path = Path(__file__).resolve().parents[2] / "data" / "estudiantes.xlsx"
        self.convocatorias = ConvocatoriaRepository()

    @property
    def supabase(self):
        """Obtiene un cliente aislado para evitar reutilizar sockets entre solicitudes."""
        return get_supabase_client()

    @property
    def usa_supabase(self) -> bool:
        return self.supabase is not None

    @staticmethod
    def _normalizar(valor: str) -> str:
        return "".join(valor.upper().split())

    def _crear_si_no_existe(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if self.path.exists():
            return
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Estudiantes"
        hoja.append(self._headers)
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = "A1:F1"
        for celda in hoja[1]:
            celda.font = celda.font.copy(bold=True)
        incidencias = libro.create_sheet("Incidencias")
        incidencias.append(self._incident_headers)
        incidencias.freeze_panes = "A2"
        incidencias.auto_filter.ref = "A1:F1"
        for celda in incidencias[1]:
            celda.font = celda.font.copy(bold=True)
        libro.save(self.path)

    def _asegurar_hoja_incidencias(self, libro) -> None:
        if "Incidencias" in libro.sheetnames:
            return
        hoja = libro.create_sheet("Incidencias")
        hoja.append(self._incident_headers)
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = "A1:F1"
        for celda in hoja[1]:
            celda.font = celda.font.copy(bold=True)

    def registrar(self, estudiante: RegistroEstudianteRequest) -> tuple[bool, str]:
        convocatoria = self.convocatorias.exigir_abierta()
        if self.usa_supabase:
            return self._registrar_supabase(estudiante, convocatoria["id"])
        with self._lock:
            self._crear_si_no_existe()
            libro = load_workbook(self.path)
            hoja = libro["Estudiantes"]
            ci = self._normalizar(estudiante.ci)
            ru = self._normalizar(estudiante.ru)
            for fila in hoja.iter_rows(min_row=2, values_only=True):
                if (ci and ci == self._normalizar(str(fila[2] or ""))) or (ru and ru == self._normalizar(str(fila[3] or ""))):
                    return False, "Tú ya te inscribiste antes. Tus datos ya están registrados."
            hoja.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), estudiante.nombre.strip(), estudiante.ci.strip(), estudiante.ru.strip(), estudiante.correo.strip(), estudiante.celular.strip()])
            libro.save(self.path)
            return True, "Datos registrados correctamente."

    def _registrar_supabase(
        self,
        estudiante: RegistroEstudianteRequest,
        convocatoria_id: str,
    ) -> tuple[bool, str]:
        try:
            existente = (
                self.supabase.table("estudiantes")
                .select("id")
                .eq("convocatoria_id", convocatoria_id)
                .or_(f"ci.eq.{estudiante.ci},ru.eq.{estudiante.ru}")
                .limit(1)
                .execute()
                .data
            )
            if existente:
                return False, "Tú ya te inscribiste antes. Tus datos ya están registrados."
            self.supabase.table("estudiantes").insert({
                "convocatoria_id": convocatoria_id,
                "nombre": estudiante.nombre.strip(), "ci": estudiante.ci.strip(), "ru": estudiante.ru.strip(),
                "correo": estudiante.correo.strip(), "celular": estudiante.celular.strip(),
            }).execute()
            return True, "Datos registrados correctamente."
        except Exception as error:
            if "duplicate" in str(error).lower() or "unique" in str(error).lower():
                return False, "Tú ya te inscribiste antes. Tus datos ya están registrados."
            raise

    def registrar_incidencia(self, nombre: str, ci: str, ru: str, archivo: str, motivo: str) -> None:
        if self.usa_supabase:
            convocatoria = self.convocatorias.obtener_actual()
            self.supabase.table("incidencias").insert({
                "convocatoria_id": convocatoria["id"] if convocatoria else None,
                "nombre": nombre.strip(), "ci": ci.strip(), "ru": ru.strip(),
                "archivo": archivo.strip(), "motivo": motivo.strip(),
            }).execute()
            return
        with self._lock:
            self._crear_si_no_existe()
            libro = load_workbook(self.path)
            self._asegurar_hoja_incidencias(libro)
            libro["Incidencias"].append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), nombre.strip(), ci.strip(), ru.strip(), archivo.strip(), motivo.strip()])
            libro.save(self.path)

    def esta_registrado(self, ci: str, ru: str) -> bool:
        if self.usa_supabase:
            convocatoria = self.convocatorias.obtener_actual()
            if convocatoria is None:
                return False
            # Evita comparar identificadores vacíos cuando el PDF usa otra etiqueta.
            for columna, valor in (("ci", ci.strip()), ("ru", ru.strip())):
                if not valor:
                    continue
                encontrado = (
                    self.supabase.table("estudiantes")
                    .select("id")
                    .eq("convocatoria_id", convocatoria["id"])
                    .eq(columna, valor)
                    .limit(1)
                    .execute()
                    .data
                )
                if encontrado:
                    return True
            return False
        with self._lock:
            if not self.path.exists():
                return False
            libro = load_workbook(self.path, read_only=True)
            for fila in libro["Estudiantes"].iter_rows(min_row=2, values_only=True):
                if (ci and self._normalizar(ci) == self._normalizar(str(fila[2] or ""))) or (ru and self._normalizar(ru) == self._normalizar(str(fila[3] or ""))):
                    return True
            return False

    def listar(self, convocatoria_id: str | None = None) -> list[dict[str, str]]:
        if self.usa_supabase:
            if convocatoria_id is None:
                convocatoria = self.convocatorias.obtener_actual()
                convocatoria_id = convocatoria["id"] if convocatoria else None
            if convocatoria_id is None:
                return []
            filas = (
                self.supabase.table("estudiantes")
                .select("created_at,nombre,ci,ru,correo,celular")
                .eq("convocatoria_id", convocatoria_id)
                .order("created_at", desc=True)
                .execute()
                .data
                or []
            )
            return [{"Fecha": f.get("created_at", ""), "Nombre": f.get("nombre", ""), "CI": f.get("ci", ""), "RU": f.get("ru", ""), "Correo": f.get("correo", ""), "Celular": f.get("celular", "")} for f in filas]
        with self._lock:
            if not self.path.exists():
                return []
            libro = load_workbook(self.path, read_only=True, data_only=True)
            return [dict(zip(self._headers, [str(valor or "") for valor in fila])) for fila in libro["Estudiantes"].iter_rows(min_row=2, values_only=True) if any(valor is not None and str(valor).strip() for valor in fila)]

    def listar_incidencias(
        self,
        convocatoria_id: str | None = None,
    ) -> list[dict[str, str]]:
        if self.usa_supabase:
            if convocatoria_id is None:
                convocatoria = self.convocatorias.obtener_actual()
                convocatoria_id = convocatoria["id"] if convocatoria else None
            if convocatoria_id is None:
                return []
            filas = (
                self.supabase.table("incidencias")
                .select("created_at,nombre,ci,ru,archivo,motivo")
                .eq("convocatoria_id", convocatoria_id)
                .order("created_at", desc=True)
                .execute()
                .data
                or []
            )
            return [{"Fecha": f.get("created_at", ""), "Nombre": f.get("nombre", ""), "CI": f.get("ci", ""), "RU": f.get("ru", ""), "Archivo": f.get("archivo", ""), "Motivo": f.get("motivo", "")} for f in filas]
        with self._lock:
            if not self.path.exists():
                return []
            libro = load_workbook(self.path, read_only=True, data_only=True)
            if "Incidencias" not in libro.sheetnames:
                return []
            return [dict(zip(self._incident_headers, [str(valor or "") for valor in fila])) for fila in libro["Incidencias"].iter_rows(min_row=2, values_only=True) if any(valor is not None and str(valor).strip() for valor in fila)]

    def generar_excel(self, convocatoria_id: str | None = None) -> bytes:
        """Genera en memoria el reporte de una convocatoria."""
        if not self.usa_supabase:
            self._crear_si_no_existe()
            libro = load_workbook(self.path)
        else:
            libro = Workbook()
        estudiantes = libro.active
        estudiantes.title = "Estudiantes"
        if self.usa_supabase:
            estudiantes.append(self._headers)
            for fila in self.listar(convocatoria_id):
                estudiantes.append([fila[campo] for campo in self._headers])
            incidencias = libro.create_sheet("Incidencias")
            incidencias.append(self._incident_headers)
            for fila in self.listar_incidencias(convocatoria_id):
                incidencias.append(
                    [fila[campo] for campo in self._incident_headers]
                )
        salida = BytesIO()
        libro.save(salida)
        salida.seek(0)
        return salida.getvalue()
