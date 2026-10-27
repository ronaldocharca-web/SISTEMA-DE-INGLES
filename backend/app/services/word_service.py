from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_border(cell, color: str = "B7B7B7", size: str = "6") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _format_cell(cell, text: str, *, bold: bool = False, color: str = "000000", align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(str(text or ""))
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_border(cell)


def generar_word_estudiantes(estudiantes: list[dict[str, str]], convocatoria: dict | None) -> bytes:
    """Genera un acta Word con una fila dinámica por estudiante registrado."""
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    datos_convocatoria = convocatoria or {}
    gestion = datos_convocatoria.get("gestion", "")
    semestre = str(datos_convocatoria.get("semestre", "")).upper()
    periodo = f"GESTIÓN {gestion}, {semestre}" if gestion and semestre else ""
    titulo = "EXAMEN DE CLASIFICACIÓN DE INGLÉS"
    if periodo:
        titulo = f"{titulo} - {periodo}"
    title_run = title.add_run(titulo)
    title_run.font.name = "Arial"
    title_run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_run.font.size = Pt(14)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle_run = subtitle.add_run(datos_convocatoria.get("nombre", "Registro de estudiantes"))
    subtitle_run.font.name = "Arial"
    subtitle_run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    subtitle_run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.bold = True
    subtitle_run.font.color.rgb = RGBColor(80, 80, 80)

    headers = ["N°", "NOMBRE", "CÉDULA", "NOTA", "NIVEL", "OBS"]
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.35, 3.05, 1.1, 0.65, 0.85, 1.2]
    header_row = table.rows[0]
    _set_repeat_table_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        cell.width = Inches(widths[index])
        _set_cell_shading(cell, "FFFFFF")
        _format_cell(cell, header, bold=True, color="000000", align=WD_ALIGN_PARAGRAPH.CENTER)

    for number, student in enumerate(estudiantes, start=1):
        values = [number, student.get("Nombre", ""), student.get("CI", ""), "", "", ""]
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            _set_cell_shading(cell, "FFFFFF")
            alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 1 else WD_ALIGN_PARAGRAPH.CENTER
            _format_cell(cell, value, align=alignment)

    if not estudiantes:
        row = table.add_row()
        merged = row.cells[0]
        for cell in row.cells[1:]:
            merged = merged.merge(cell)
        _set_cell_shading(merged, "FFFFFF")
        _format_cell(merged, "No hay estudiantes registrados en esta convocatoria.", color="666666", align=WD_ALIGN_PARAGRAPH.CENTER)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output.getvalue()
